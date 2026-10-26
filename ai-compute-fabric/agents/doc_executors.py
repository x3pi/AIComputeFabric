"""DOCUMENT_PARSE and OCR executors (plan P5.1) on local engines: pypdf / python-docx / stdlib for documents,
the tesseract CLI for images.

Inputs name files by a reference RELATIVE to a data root; anything that escapes the root (absolute path, `..`,
symlink out) is refused, as are files over the size limits. Not covered: scanned PDFs (a PDF page with no text layer
returns an empty page, it is not OCR'd), PDF table extraction, handwriting, layout analysis, Windows."""
from __future__ import annotations

import subprocess
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from agents.tasks import Task, TaskResult, TaskType, parse_input

MAX_FILE = 50 * 1024 * 1024
MAX_UNZIPPED = 200 * 1024 * 1024
MAX_PAGES = 2000
MAX_PIXELS = 40_000_000
MIMES = {"application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "text/plain", "text/markdown", "text/html"}


class RefusedInput(Exception):
    """The reference or file is not acceptable (path escape, too big, wrong type). Retrying will not help."""


def resolve(root: Path, ref: str, limit: int | None = None) -> Path:
    limit = MAX_FILE if limit is None else limit
    if not ref or "\x00" in ref:
        raise RefusedInput("empty or malformed reference")
    p = Path(ref)
    if p.is_absolute() or ".." in p.parts:
        raise RefusedInput("reference must be relative and stay inside the data root")
    full = (root / p).resolve()
    if not full.is_relative_to(root.resolve()):
        raise RefusedInput("reference resolves outside the data root")
    if not full.is_file():
        raise RefusedInput("no such file")
    if full.stat().st_size > limit:
        raise RefusedInput(f"file larger than {limit} bytes")
    return full


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        elif tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3"):
            self.out.append("\n")

    def handle_data(self, d):
        if not self.skip:
            self.out.append(d)


def _check_zip(path: Path):
    try:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise RefusedInput("document expands beyond the size limit")
    except zipfile.BadZipFile:
        raise RefusedInput("not a valid docx") from None


class DocumentExecutor:
    def __init__(self, root: str | Path):
        self.root = Path(root)

    def parse(self, ref: str, mime: str) -> dict:
        if mime not in MIMES:
            raise RefusedInput(f"unsupported mime {mime!r}")
        path = resolve(self.root, ref)
        if mime == "application/pdf":
            from pypdf import PdfReader
            from pypdf.errors import PyPdfError
            try:
                r = PdfReader(str(path))
                if r.is_encrypted:
                    raise RefusedInput("encrypted PDF")
                if len(r.pages) > MAX_PAGES:
                    raise RefusedInput("too many pages")
                return {"pages": [(pg.extract_text() or "").strip() for pg in r.pages], "tables": []}
            except PyPdfError as e:
                raise RefusedInput(f"unreadable PDF: {e}") from None
        if mime.endswith("wordprocessingml.document"):
            _check_zip(path)
            import docx
            d = docx.Document(str(path))
            tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
            return {"pages": ["\n".join(p.text for p in d.paragraphs).strip()], "tables": tables}
        raw = path.read_bytes().decode("utf-8", errors="replace")
        if mime == "text/html":
            p = _Text()
            p.feed(raw)
            raw = "".join(p.out)
        return {"pages": [raw.strip()], "tables": []}

    def execute(self, task: Task, attempt: int = 1, deps: dict | None = None) -> TaskResult:
        if task.type != TaskType.DOCUMENT_PARSE:
            raise NotImplementedError(task.type.value)
        inp = parse_input(task)
        return TaskResult(task_id=task.task_id, type=task.type, output=self.parse(inp.document_ref, inp.mime), worker_id="docparse")


class OcrExecutor:
    """tesseract CLI. `lang` is mapped to tesseract codes; confidence is the mean word confidence / 100."""
    LANGS = {"vi": "vie", "en": "eng", "vi+en": "vie+eng"}

    def __init__(self, root: str | Path, tesseract: str = "tesseract", timeout: float = 60.0):
        self.root, self.tesseract, self.timeout = Path(root), tesseract, timeout

    def read(self, ref: str, lang: str) -> dict:
        if lang not in self.LANGS:
            raise RefusedInput(f"unsupported language {lang!r}")
        path = resolve(self.root, ref, limit=30 * 1024 * 1024)
        from PIL import Image, UnidentifiedImageError
        try:
            with Image.open(path) as im:
                if im.width * im.height > MAX_PIXELS:
                    raise RefusedInput("image too large")
                im.verify()
        except (UnidentifiedImageError, OSError):
            raise RefusedInput("not a readable image") from None
        try:
            r = subprocess.run([self.tesseract, str(path), "stdout", "-l", self.LANGS[lang], "--psm", "6", "tsv"],
                               capture_output=True, text=True, timeout=self.timeout)
        except subprocess.TimeoutExpired:
            raise RefusedInput("OCR timed out") from None
        if r.returncode != 0:
            raise RuntimeError(f"tesseract failed: {r.stderr[:200]}")
        lines: dict[tuple, list] = {}
        confs = []
        for row in r.stdout.splitlines()[1:]:
            f = row.split("\t")
            if len(f) < 12 or not f[11].strip():
                continue
            c = float(f[10])
            if c < 0:
                continue
            lines.setdefault((f[2], f[3], f[4]), []).append(f[11])
            confs.append(c)
        return {"text": "\n".join(" ".join(w) for w in lines.values()), "confidence": round(sum(confs) / len(confs) / 100, 4) if confs else 0.0}

    def execute(self, task: Task, attempt: int = 1, deps: dict | None = None) -> TaskResult:
        if task.type != TaskType.OCR:
            raise NotImplementedError(task.type.value)
        inp = parse_input(task)
        return TaskResult(task_id=task.task_id, type=task.type, output=self.read(inp.image_ref, inp.lang), worker_id="tesseract")
