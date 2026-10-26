"""DOCUMENT_PARSE (pypdf / python-docx / stdlib) and OCR (tesseract + vie) executors on files generated here.
OCR accuracy is measured on SYNTHETIC rendered text (clean and degraded), not on real scans; the numbers say what
the engine does on that input, not on a customer's scanned invoices. Skipped when tesseract/Vietnamese data is missing."""
import io
import json
import os
import random
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from agents.doc_executors import DocumentExecutor, OcrExecutor, RefusedInput
from agents.tasks import Task, TaskType
from agents.validator import validate_result
from audio.evaluation.vn_normalizer import cer

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
HAVE_TESS = bool(shutil.which("tesseract")) and "vie" in subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True).stdout
needs_font = pytest.mark.skipif(not Path(FONT).exists(), reason="DejaVu font missing")
needs_tess = pytest.mark.skipif(not HAVE_TESS, reason="tesseract with Vietnamese data not installed")
LINES = ["Hóa đơn số 0001234 ngày 15/03/2025", "Công ty TNHH Thương mại Việt Phát", "Tổng tiền thanh toán: 12.500.000 đồng",
         "Địa chỉ: 25 Nguyễn Huệ, Quận 1, Thành phố Hồ Chí Minh", "Cảm ơn quý khách đã sử dụng dịch vụ"]


def task(tp, inp):
    return Task(task_id="t", tenant_id="acme", type=tp, input=inp)


def make_pdf(path, pages):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(TTFont("DejaVu", FONT))
    c = canvas.Canvas(str(path), pagesize=A4)
    for t in pages:
        c.setFont("DejaVu", 14)
        c.drawString(60, 760, t)
        c.showPage()
    c.save()


def render(lines, size=32, blur=0.0, noise=0, rot=0.0):
    f = ImageFont.truetype(FONT, size)
    im = Image.new("L", (1500, int(size * 1.9 * len(lines)) + 40), 255)
    d = ImageDraw.Draw(im)
    for i, t in enumerate(lines):
        d.text((20, 20 + i * size * 1.9), t, font=f, fill=0)
    if rot:
        im = im.rotate(rot, expand=False, fillcolor=255)
    if blur:
        im = im.filter(ImageFilter.GaussianBlur(blur))
    if noise:
        rng = random.Random(1)
        px = im.load()
        for _ in range(im.width * im.height // 12):
            x, y = rng.randrange(im.width), rng.randrange(im.height)
            px[x, y] = max(0, min(255, px[x, y] + rng.randint(-noise, noise)))
    return im


@pytest.fixture
def root(tmp_path):
    return tmp_path


# ---------------------------------------------------------------- DOCUMENT_PARSE
@needs_font
def test_pdf_pages_keep_vietnamese_text(root):
    make_pdf(root / "a.pdf", ["Hóa đơn số 1", "Tổng tiền 12.500.000 đồng", "Trang ba"])
    t = task(TaskType.DOCUMENT_PARSE, {"document_ref": "a.pdf", "mime": "application/pdf"})
    res = DocumentExecutor(root).execute(t)
    assert res.output["pages"] == ["Hóa đơn số 1", "Tổng tiền 12.500.000 đồng", "Trang ba"]
    assert validate_result(t, res).ok


def test_docx_paragraphs_and_tables(root):
    import docx
    d = docx.Document()
    d.add_paragraph("Hợp đồng cung cấp dịch vụ")
    d.add_paragraph("Bên A: Công ty Việt Phát")
    tb = d.add_table(rows=2, cols=2)
    for i, row in enumerate([["Hạng mục", "Số tiền"], ["Phí tháng 3", "1.000.000"]]):
        for j, v in enumerate(row):
            tb.cell(i, j).text = v
    d.save(root / "c.docx")
    t = task(TaskType.DOCUMENT_PARSE, {"document_ref": "c.docx", "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
    res = DocumentExecutor(root).execute(t)
    assert "Hợp đồng cung cấp dịch vụ" in res.output["pages"][0] and "Bên A" in res.output["pages"][0]
    assert res.output["tables"] == [[["Hạng mục", "Số tiền"], ["Phí tháng 3", "1.000.000"]]]
    assert validate_result(t, res).ok


def test_html_drops_scripts_and_styles(root):
    (root / "p.html").write_text("<html><style>p{color:red}</style><script>alert('x')</script><p>Xin chào</p><div>thế giới</div></html>", encoding="utf-8")
    out = DocumentExecutor(root).parse("p.html", "text/html")
    assert "Xin chào" in out["pages"][0] and "thế giới" in out["pages"][0]
    assert "alert" not in out["pages"][0] and "color" not in out["pages"][0]


def test_plain_text_with_bad_bytes_does_not_crash(root):
    (root / "t.txt").write_bytes("Tiếng Việt".encode() + b"\xff\xfe broken")
    assert "Tiếng Việt" in DocumentExecutor(root).parse("t.txt", "text/plain")["pages"][0]


@pytest.mark.parametrize("ref", ["../x.txt", "/etc/passwd", "a/../../x.txt", "", "a\x00.txt", "missing.txt"])
def test_references_that_escape_or_do_not_exist_are_refused(root, ref):
    (root.parent / "x.txt").write_text("secret")
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse(ref, "text/plain")


def test_symlink_pointing_outside_the_root_is_refused(root, tmp_path_factory):
    outside = tmp_path_factory.mktemp("outside")
    (outside / "s.txt").write_text("secret")
    os.symlink(outside / "s.txt", root / "link.txt")
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("link.txt", "text/plain")


def test_wrong_mime_oversize_corrupt_and_zip_bomb_are_refused(root, monkeypatch):
    (root / "t.txt").write_text("x")
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("t.txt", "application/x-msdownload")
    (root / "bad.pdf").write_bytes(b"%PDF-1.4 this is not a pdf")
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("bad.pdf", "application/pdf")
    (root / "bad.docx").write_bytes(b"not a zip")
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("bad.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (210 * 1024 * 1024))   # 210 MB of zeros compresses to ~200 KB
    (root / "bomb.docx").write_bytes(buf.getvalue())
    assert (root / "bomb.docx").stat().st_size < 5 * 1024 * 1024
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("bomb.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    monkeypatch.setattr("agents.doc_executors.MAX_FILE", 10)
    (root / "big.txt").write_text("x" * 100)
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("big.txt", "text/plain")


@needs_font
def test_encrypted_pdf_is_refused(root):
    from pypdf import PdfReader, PdfWriter
    make_pdf(root / "plain.pdf", ["bí mật"])
    w = PdfWriter()
    for p in PdfReader(str(root / "plain.pdf")).pages:
        w.add_page(p)
    w.encrypt("pw")
    with open(root / "enc.pdf", "wb") as f:
        w.write(f)
    with pytest.raises(RefusedInput):
        DocumentExecutor(root).parse("enc.pdf", "application/pdf")


# ---------------------------------------------------------------- OCR
@needs_font
@needs_tess
def test_ocr_reads_clean_vietnamese_text_with_low_error(root):
    render(LINES).save(root / "clean.png")
    t = task(TaskType.OCR, {"image_ref": "clean.png", "lang": "vi"})
    res = OcrExecutor(root).execute(t)
    err = cer(" ".join(LINES), res.output["text"])
    assert err < 0.03, (err, res.output["text"])
    assert res.output["confidence"] > 0.8
    assert validate_result(t, res).ok


@needs_font
@needs_tess
def test_ocr_degraded_input_is_worse_but_reported_not_hidden(root):
    results = {}
    for name, im in {"clean": render(LINES), "small": render(LINES, size=11), "blur": render(LINES, blur=2.2), "noise": render(LINES, noise=230), "tilt": render(LINES, rot=3)}.items():
        im.save(root / f"{name}.png")
        out = OcrExecutor(root).read(f"{name}.png", "vi")
        results[name] = (cer(" ".join(LINES), out["text"]), out["confidence"])
    print({k: (round(v[0], 3), round(v[1], 2)) for k, v in results.items()})
    Path(os.environ.get("AICF_OCR_REPORT", os.devnull)).write_text(json.dumps(
        {"set": "5 Vietnamese invoice-like lines rendered with DejaVu Sans, 1500px wide; degradations applied synthetically",
         "cer_and_mean_confidence": {k: {"cer": round(v[0], 4), "confidence": v[1]} for k, v in results.items()},
         "caveat": "synthetic rendering, one font, 5 lines: shows the engine's behaviour under controlled degradation, not accuracy on real scans or photos"}, indent=1))
    assert results["clean"][0] < 0.03
    assert results["blur"][0] > results["clean"][0] or results["noise"][0] > results["clean"][0]   # degradation is visible, not masked
    assert results["noise"][1] < results["clean"][1]                                              # and the confidence goes down with it


@needs_tess
def test_ocr_refusals(root):
    (root / "t.png").write_text("not an image")
    with pytest.raises(RefusedInput):
        OcrExecutor(root).read("t.png", "vi")
    Image.new("L", (10, 10), 255).save(root / "ok.png")
    with pytest.raises(RefusedInput):
        OcrExecutor(root).read("ok.png", "klingon")
    with pytest.raises(RefusedInput):
        OcrExecutor(root).read("../ok.png", "vi")
    Image.new("1", (7000, 7000), 1).save(root / "huge.png")
    with pytest.raises(RefusedInput):
        OcrExecutor(root).read("huge.png", "vi")


@needs_tess
def test_blank_image_gives_empty_text_and_zero_confidence(root):
    Image.new("L", (400, 200), 255).save(root / "blank.png")
    out = OcrExecutor(root).read("blank.png", "vi")
    assert out["text"].strip() == "" and out["confidence"] == 0.0
